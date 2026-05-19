from pathlib import Path

from docx import Document


def build_docx(path: Path, pack: dict):
    doc = Document()
    doc.add_heading("AI Admin Assistant Output", level=1)

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(pack.get("summary", ""))

    doc.add_heading("Action Tasks", level=2)
    for item in pack.get("tasks", []):
        doc.add_paragraph(
            f"[{item.get('priority', 'Medium')}] {item.get('task', '')} (Owner: {item.get('owner', 'You')})",
            style="List Bullet",
        )

    doc.add_heading("Draft Emails", level=2)
    for draft in pack.get("email_drafts", []):
        doc.add_paragraph(f"Subject: {draft.get('subject', '')}")
        doc.add_paragraph(draft.get("body", ""))

    doc.add_heading("Reminder Suggestions", level=2)
    for rem in pack.get("reminders", []):
        doc.add_paragraph(f"{rem.get('when', '')} - {rem.get('text', '')}", style="List Bullet")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
